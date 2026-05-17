from flask import Flask, render_template, request, redirect, session, flash, send_from_directory
from app.api.middleware import login_required, admin_required
from app.core.database import get_db
from app.core.logs import log_operacao
from app.config import Config
from datetime import datetime
import os
import re

def _senha_forte(senha):
    return (
        len(senha) >= 8 and
        re.search(r'[a-z]', senha) and
        re.search(r'[A-Z]', senha) and
        re.search(r'\d', senha) and
        re.search(r'[\W_]', senha)
    )

_MSG_SENHA = "A senha deve ter no mínimo 8 caracteres, incluindo letras maiúsculas, minúsculas, números e caracteres especiais."

def register_routes(app):
    """Registra todas as rotas web"""
    
    # ==================== AUTENTICAÇÃO ====================
    @app.route("/", methods=["GET", "POST"])
    def login():
        if request.method == "POST":
            with get_db() as conn:
                user = conn.execute(
                    "SELECT id, role, username FROM usuarios WHERE username=? AND password=?",
                    (request.form["u"], request.form["s"])
                ).fetchone()
            if user:
                session['user_id'] = user['id']
                session['role'] = user['role']
                session['username'] = user['username']
                return redirect("/inicio")
        return render_template("auth/login.html")
    
    @app.route("/inicio")
    @login_required
    def inicio():
        return render_template("dashboard/home.html")
    
    @app.route("/logout")
    def logout():
        session.clear()
        return redirect("/")
    
    # ==================== ALUNOS ====================
    @app.route("/cadastro_aluno", methods=["GET", "POST"])
    @login_required
    def cadastro_aluno():
        from app.repositories.students import StudentRepository
        repo = StudentRepository()
        
        if request.method == "POST":
            nome = request.form.get("nome", "").strip()
            turma = request.form.get("turma", "").strip()
            serie = request.form.get("serie", "").strip()
            
            if not nome or not turma or not serie:
                flash("Preencher nome, turma e série", "error")
            else:
                # Buscar horários padrão
                with get_db() as conn:
                    padrao = conn.execute(
                        "SELECT saida_seg, saida_ter, saida_qua, saida_qui, saida_sex FROM horarios_padrao WHERE serie = ?",
                        (serie,)
                    ).fetchone()
                
                data = {
                    'nome': nome, 'turma': turma, 'serie': serie,
                    'saida_seg': request.form.get("saida_seg") or (padrao['saida_seg'] if padrao else ''),
                    'saida_ter': request.form.get("saida_ter") or (padrao['saida_ter'] if padrao else ''),
                    'saida_qua': request.form.get("saida_qua") or (padrao['saida_qua'] if padrao else ''),
                    'saida_qui': request.form.get("saida_qui") or (padrao['saida_qui'] if padrao else ''),
                    'saida_sex': request.form.get("saida_sex") or (padrao['saida_sex'] if padrao else ''),
                    'responsaveis': request.form.get("responsaveis", ""),
                    'telefone': request.form.get("telefone", ""),
                    'email_responsavel': request.form.get("email_responsavel", ""),
                    'data_nascimento': request.form.get("data_nascimento", ""),
                    'alergias': request.form.get("alergias", ""),
                    'observacoes': request.form.get("observacoes", ""),
                    'foto_path': None
                }
                
                # Upload de foto
                if 'foto' in request.files:
                    file = request.files['foto']
                    if file and file.filename != '':
                        from app.core.validators import allowed_file
                        from werkzeug.utils import secure_filename
                        if allowed_file(file.filename):
                            filename = secure_filename(f"{datetime.now().strftime('%Y%m%d%H%M%S')}_{file.filename}")
                            file.save(os.path.join(app.config['UPLOAD_FOLDER'], 'photos', filename))
                            data['foto_path'] = os.path.join('photos', filename)
                
                aluno_id = repo.create(data)
                log_operacao(session.get('username'), "CADASTROU ALUNO", f"Nome: {nome}")
                flash(f"Aluno {nome} cadastrado!", "success")
                return redirect("/cadastro_aluno")
        
        # GET - mostrar lista
        busca = request.args.get("busca")
        if busca:
            alunos = repo.search_by_name(busca)
        else:
            alunos = repo.get_all_grouped()
        
        return render_template("students/list_of_students.html", alunos=alunos, busca=busca, series=Config.SERIES)
    
    # ==================== SAÍDAS ====================
    @app.route("/saidas")
    @login_required
    def lista_saidas():
        from app.repositories.departures import DepartureRepository
        repo = DepartureRepository()
        
        data_selecionada = request.args.get("data", datetime.now().strftime("%Y-%m-%d"))
        busca = request.args.get("busca", "").strip()
        
        # Limpar saídas antigas
        repo.cleanup_old(30)
        
        saidas = repo.get_by_date(data_selecionada, busca)
        
        pendentes = [s for s in saidas if s['status'] == 'pendente']
        concluidas = [s for s in saidas if s['status'] == 'concluida']
        
        return render_template("departures/list_of_exits.html",
                               pendentes=pendentes,
                               concluidas=concluidas,
                               data_selecionada=data_selecionada,
                               busca=busca)
    
    @app.route("/concluir_saida/<int:id_saida>")
    @login_required
    def concluir_saida(id_saida):
        from app.repositories.departures import DepartureRepository
        from app.core.eamail import enviar_email
        repo = DepartureRepository()
        
        with get_db() as conn:
            info = conn.execute("""
                SELECT s.aluno, s.data_saida, s.horario, s.motivo, s.responsavel_escola,
                       s.tipo_saida, s.acompanhante, a.email_responsavel, a.nome
                FROM saidas s
                JOIN alunos a ON s.aluno = a.id
                WHERE s.id = ?
            """, (id_saida,)).fetchone()
        
        if info:
            repo.complete(id_saida, session['user_id'])
            
            if info['email_responsavel']:
                corpo = f"""
                <h2>Saída autorizada</h2>
                <p>O aluno <strong>{info['nome']}</strong> teve a saída antecipada autorizada.</p>
                <ul>
                    <li>Data: {info['data_saida']}</li>
                    <li>Horário: {info['horario']}</li>
                    <li>Motivo: {info['motivo']}</li>
                    <li>Responsável na escola: {info['responsavel_escola']}</li>
                </ul>
                """
                enviar_email(info['email_responsavel'], f"Saída autorizada - {info['nome']}", corpo)
            
            log_operacao(session.get('username'), "CONCLUIU SAÍDA", f"ID Saída: {id_saida}")
            flash("Saída autorizada!", "success")
        
        return redirect("/saidas")
    
    # ==================== ADMIN ====================
    @app.route("/configuracoes")
    @admin_required
    def configuracoes():
        from app.repositories.students import StudentRepository
        from app.repositories.users import UserRepository
        from app.repositories.departures import DepartureRepository
        import glob
        
        student_repo = StudentRepository()
        user_repo = UserRepository()
        departure_repo = DepartureRepository()
        
        total_alunos = student_repo.count()
        total_usuarios = user_repo.count()
        total_saidas = departure_repo.count()
        
        # Saídas de hoje
        hoje = datetime.now().strftime("%Y-%m-%d")
        saidas_hoje = len(departure_repo.get_by_date(hoje))
        
        # Backups
        backups = glob.glob("backups/*.db")
        total_backups = len(backups)
        ultimo_backup = backups[-1].split('/')[-1] if backups else None
        
        return render_template("admin/settings.html",
                               total_alunos=total_alunos,
                               total_usuarios=total_usuarios,
                               total_saidas=total_saidas,
                               saidas_hoje=saidas_hoje,
                               total_backups=total_backups,
                               ultimo_backup=ultimo_backup)
    
    @app.route("/admin/backup")
    @admin_required
    def admin_backup():
        from app.core.backup import backup_database
        try:
            backup_path = backup_database()
            return redirect(f"/configuracoes?backup=ok&arquivo={backup_path.split('/')[-1]}")
        except Exception as e:
            return redirect(f"/configuracoes?backup=erro&msg={str(e)}")
    
    # ==================== ARQUIVOS ESTÁTICOS ====================
    @app.route("/uploads/<path:filename>")
    @login_required
    def uploaded_file(filename):
        return send_from_directory(app.config['UPLOAD_FOLDER'], filename)
    
    # ==================== MANUAL ====================
    @app.route("/manual")
    @login_required
    def manual():
        return render_template("help/manual.html")

    # ==================== EDITAR ALUNO ====================
    @app.route("/editar_aluno/<int:id>", methods=["GET", "POST"])
    @login_required
    def editar_aluno(id):
        from app.repositories.students import StudentRepository
        repo = StudentRepository()
        aluno = repo.get_by_id(id)
        if not aluno:
            flash("Aluno não encontrado.", "error")
            return redirect("/cadastro_aluno")
        if request.method == "POST":
            data = {
                'nome': request.form.get('nome', '').strip(),
                'turma': request.form.get('turma', '').strip(),
                'serie': request.form.get('serie', '').strip(),
                'responsaveis': request.form.get('responsaveis', ''),
                'telefone': request.form.get('telefone', ''),
                'email_responsavel': request.form.get('email_responsavel', ''),
                'data_nascimento': request.form.get('data_nascimento', ''),
                'alergias': request.form.get('alergias', ''),
                'observacoes': request.form.get('observacoes', ''),
                'saida_seg': request.form.get('saida_seg', ''),
                'saida_ter': request.form.get('saida_ter', ''),
                'saida_qua': request.form.get('saida_qua', ''),
                'saida_qui': request.form.get('saida_qui', ''),
                'saida_sex': request.form.get('saida_sex', ''),
                'foto_path': aluno.get('foto_path'),
            }
            if not data['nome'] or not data['turma'] or not data['serie']:
                flash("Nome, turma e série são obrigatórios.", "error")
            else:
                repo.update(id, data)
                log_operacao(session.get('username'), "EDITOU ALUNO", f"ID: {id}, Nome: {data['nome']}")
                flash(f"Aluno {data['nome']} atualizado!", "success")
                return redirect("/cadastro_aluno")
        return render_template("students/edit_students.html", aluno=aluno, series=Config.SERIES)

    # ==================== DELETAR ALUNO ====================
    @app.route("/deletar_aluno/<int:id>")
    @admin_required
    def deletar_aluno(id):
        from app.repositories.students import StudentRepository
        repo = StudentRepository()
        aluno = repo.get_by_id(id)
        if aluno:
            repo.delete(id)
            log_operacao(session.get('username'), "DELETOU ALUNO", f"ID: {id}, Nome: {aluno['nome']}")
            flash(f"Aluno {aluno['nome']} removido.", "success")
        return redirect("/cadastro_aluno")

    # ==================== HISTÓRICO DO ALUNO ====================
    @app.route("/historico_aluno/<int:id>")
    @login_required
    def historico_aluno(id):
        from app.repositories.students import StudentRepository
        from app.repositories.departures import DepartureRepository
        aluno = StudentRepository().get_by_id(id)
        historico = DepartureRepository().get_historic(id) if aluno else []
        return render_template("students/history_students.html", aluno=aluno, historico=historico)

    # ==================== REGISTRAR SAÍDA ====================
    @app.route("/registrar_saida", methods=["GET", "POST"])
    @login_required
    def registrar_saida():
        from app.repositories.students import StudentRepository
        from app.repositories.departures import DepartureRepository
        aluno_selecionado = request.args.get("aluno_id", "")
        alunos = StudentRepository().get_all_grouped()
        today = datetime.now().strftime("%Y-%m-%d")

        if request.method == "POST":
            aluno_id = request.form.get("aluno_id")
            data_saida = request.form.get("data_saida", today)
            horario = request.form.get("horario", "").strip()
            motivo = request.form.get("motivo", "").strip()
            responsavel_escola = request.form.get("responsavel_escola", "").strip()
            tipo_saida = request.form.get("tipo_saida", "").strip()
            acompanhante = request.form.get("acompanhante", "").strip() if tipo_saida == "acompanhado" else None

            if not all([aluno_id, horario, motivo, responsavel_escola, tipo_saida]):
                erro = "Preencha todos os campos obrigatórios."
                return render_template("students/register.html", alunos=alunos,
                                       aluno_selecionado=aluno_selecionado, erro=erro, today=today)

            documento_path = None
            if 'documento' in request.files:
                file = request.files['documento']
                if file and file.filename != '':
                    from app.core.validators import allowed_file
                    from werkzeug.utils import secure_filename
                    if allowed_file(file.filename):
                        filename = secure_filename(f"{datetime.now().strftime('%Y%m%d%H%M%S')}_{file.filename}")
                        file.save(os.path.join(app.config['UPLOAD_FOLDER'], 'documents', filename))
                        documento_path = os.path.join('documents', filename)

            DepartureRepository().create({
                'aluno': int(aluno_id),
                'data_saida': data_saida,
                'horario': horario,
                'motivo': motivo,
                'responsavel_escola': responsavel_escola,
                'tipo_saida': tipo_saida,
                'acompanhante': acompanhante,
                'documento_path': documento_path,
            })
            log_operacao(session.get('username'), "REGISTROU SAÍDA", f"Aluno ID: {aluno_id}")
            flash("Saída registrada com sucesso!", "success")
            return redirect("/saidas")

        return render_template("students/register.html", alunos=alunos,
                               aluno_selecionado=aluno_selecionado, erro=None, today=today)

    # ==================== NOVA SAÍDA (atalho) ====================
    @app.route("/novo")
    @login_required
    def novo():
        return redirect("/registrar_saida")

    # ==================== EDITAR SAÍDA ====================
    @app.route("/editar_saida/<int:id>", methods=["GET", "POST"])
    @login_required
    def editar_saida(id):
        from app.repositories.departures import DepartureRepository
        repo = DepartureRepository()

        with get_db() as conn:
            row = conn.execute("""
                SELECT s.*, a.nome as aluno_nome
                FROM saidas s
                JOIN alunos a ON s.aluno = a.id
                WHERE s.id = ? AND s.status = 'pendente'
            """, (id,)).fetchone()

        if not row:
            flash("Saída não encontrada ou já concluída.", "error")
            return redirect("/saidas")

        saida = dict(row)
        saida['aluno'] = saida['aluno_nome']

        if request.method == "POST":
            data = {
                'horario': request.form.get('horario', '').strip(),
                'motivo': request.form.get('motivo', '').strip(),
                'responsavel_escola': request.form.get('responsavel_escola', '').strip(),
                'tipo_saida': request.form.get('tipo_saida', '').strip(),
                'acompanhante': request.form.get('acompanhante', '').strip() or None,
            }
            repo.update(id, data)
            log_operacao(session.get('username'), "EDITOU SAÍDA", f"ID: {id}")
            flash("Saída atualizada!", "success")
            return redirect("/saidas")

        return render_template("departures/edit_exits.html", saida=saida)

    # ==================== HISTÓRICO GERAL ====================
    @app.route("/historico")
    @login_required
    def historico():
        from app.repositories.departures import DepartureRepository
        nome = request.args.get("nome", "").strip()
        resultados = DepartureRepository().get_by_filters({'nome': nome}) if nome else []
        return render_template("departures/history_of_departures.html", resultados=resultados, nome=nome)

    # ==================== ADMIN - USUÁRIOS ====================
    @app.route("/admin/usuarios", methods=["GET", "POST"])
    @admin_required
    def admin_usuarios():
        from app.repositories.users import UserRepository
        repo = UserRepository()

        if request.method == "POST":
            username = request.form.get("u", "").strip()
            password = request.form.get("s", "").strip()
            email = request.form.get("e", "").strip()
            role = request.form.get("r", "basico")

            if not username or not password or not email:
                flash("Preencha todos os campos.", "error")
            elif not _senha_forte(password):
                flash(_MSG_SENHA, "error")
            elif repo.get_by_username(username):
                flash("Nome de usuário já existe.", "error")
            else:
                repo.create({'username': username, 'password': password, 'role': role, 'email': email})
                log_operacao(session.get('username'), "CRIOU USUÁRIO", f"Username: {username}")
                flash(f"Usuário {username} criado!", "success")

        usuarios = repo.get_all_without_passwords()
        return render_template("admin/users.html", usuarios=usuarios)

    # ==================== DELETAR USUÁRIO ====================
    @app.route("/deletar_usuario/<int:id>")
    @admin_required
    def deletar_usuario(id):
        from app.repositories.users import UserRepository
        repo = UserRepository()

        if id == session.get('user_id'):
            flash("Você não pode deletar sua própria conta.", "error")
            return redirect("/admin/usuarios")

        if repo.get_admin_count_excluding(id) == 0:
            flash("Não é possível deletar o único administrador.", "error")
            return redirect("/admin/usuarios")

        user = repo.get_by_id(id)
        if user:
            repo.delete(id)
            log_operacao(session.get('username'), "DELETOU USUÁRIO", f"ID: {id}, Username: {user['username']}")
            flash(f"Usuário {user['username']} removido.", "success")
        return redirect("/admin/usuarios")

    # ==================== HORÁRIOS PADRÃO ====================
    @app.route("/configurar_horarios", methods=["GET", "POST"])
    @admin_required
    def configurar_horarios():
        if request.method == "POST":
            with get_db() as conn:
                for serie in Config.SERIES:
                    conn.execute("""
                        INSERT INTO horarios_padrao (serie, saida_seg, saida_ter, saida_qua, saida_qui, saida_sex)
                        VALUES (?, ?, ?, ?, ?, ?)
                        ON CONFLICT(serie) DO UPDATE SET
                            saida_seg=excluded.saida_seg,
                            saida_ter=excluded.saida_ter,
                            saida_qua=excluded.saida_qua,
                            saida_qui=excluded.saida_qui,
                            saida_sex=excluded.saida_sex
                    """, (
                        serie,
                        request.form.get(f"seg_{serie}", ""),
                        request.form.get(f"ter_{serie}", ""),
                        request.form.get(f"qua_{serie}", ""),
                        request.form.get(f"qui_{serie}", ""),
                        request.form.get(f"sex_{serie}", ""),
                    ))
            log_operacao(session.get('username'), "ATUALIZOU HORÁRIOS", "")
            flash("Horários salvos!", "success")
            return redirect("/configurar_horarios")

        with get_db() as conn:
            rows = conn.execute("SELECT * FROM horarios_padrao").fetchall()
        horarios = {
            dict(row)['serie']: {
                'seg': dict(row).get('saida_seg', ''),
                'ter': dict(row).get('saida_ter', ''),
                'qua': dict(row).get('saida_qua', ''),
                'qui': dict(row).get('saida_qui', ''),
                'sex': dict(row).get('saida_sex', ''),
            }
            for row in rows
        }
        return render_template("admin/schedules.html", horarios=horarios, series=Config.SERIES)

    # ==================== CADASTRO EM MASSA ====================
    @app.route("/cadastro_massa", methods=["GET", "POST"])
    @admin_required
    def cadastro_massa():
        from app.repositories.students import StudentRepository

        with get_db() as conn:
            rows = conn.execute("SELECT * FROM horarios_padrao").fetchall()
        horarios_padrao = {
            dict(row)['serie']: {
                'seg': dict(row).get('saida_seg', ''),
                'ter': dict(row).get('saida_ter', ''),
                'qua': dict(row).get('saida_qua', ''),
                'qui': dict(row).get('saida_qui', ''),
                'sex': dict(row).get('saida_sex', ''),
            }
            for row in rows
        }

        erro = None

        if request.method == "POST":
            # Import via Excel
            if 'arquivo_excel' in request.files and request.files['arquivo_excel'].filename:
                file = request.files['arquivo_excel']
                try:
                    import openpyxl
                    import tempfile
                    from werkzeug.utils import secure_filename

                    with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
                        file.save(tmp.name)
                        tmp_path = tmp.name

                    wb = openpyxl.load_workbook(tmp_path)
                    ws = wb.active
                    headers = [str(c.value).strip().lower() if c.value else '' for c in ws[1]]

                    fotos_map = {}
                    if 'fotos' in request.files:
                        for foto in request.files.getlist('fotos'):
                            if foto.filename:
                                nome_sem_ext = os.path.splitext(foto.filename)[0].strip()
                                fn = secure_filename(f"{datetime.now().strftime('%Y%m%d%H%M%S')}_{foto.filename}")
                                foto.save(os.path.join(app.config['UPLOAD_FOLDER'], 'photos', fn))
                                fotos_map[nome_sem_ext.lower()] = os.path.join('photos', fn)

                    repo = StudentRepository()
                    count = 0
                    erros = []
                    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                        row_data = dict(zip(headers, row))
                        nome = str(row_data.get('nome', '') or '').strip()
                        turma = str(row_data.get('turma', '') or '').strip()
                        serie = str(row_data.get('serie', '') or '').strip()
                        if not nome or not turma or not serie:
                            erros.append(f"Linha {i}: dados obrigatórios faltando")
                            continue
                        repo.create({
                            'nome': nome, 'turma': turma, 'serie': serie,
                            'saida_seg': str(row_data.get('saida_seg', '') or ''),
                            'saida_ter': str(row_data.get('saida_ter', '') or ''),
                            'saida_qua': str(row_data.get('saida_qua', '') or ''),
                            'saida_qui': str(row_data.get('saida_qui', '') or ''),
                            'saida_sex': str(row_data.get('saida_sex', '') or ''),
                            'responsaveis': str(row_data.get('responsaveis', '') or ''),
                            'telefone': '', 'email_responsavel': '',
                            'data_nascimento': '', 'alergias': '', 'observacoes': '',
                            'foto_path': fotos_map.get(nome.lower()),
                        })
                        count += 1

                    os.unlink(tmp_path)
                    msg = f"{count} aluno(s) importado(s)."
                    if erros:
                        msg += f" {len(erros)} linha(s) ignorada(s): " + "; ".join(erros[:3])
                    log_operacao(session.get('username'), "IMPORTOU ALUNOS", msg)
                    flash(msg, "success" if not erros else "warning")
                    return redirect("/cadastro_aluno")

                except ImportError:
                    erro = "Para importar Excel instale: <code>pip3 install openpyxl</code>"
                except Exception as e:
                    erro = f"Erro ao processar arquivo: {str(e)}"

            else:
                # Individual registration
                nome = request.form.get("nome", "").strip()
                turma = request.form.get("turma", "").strip()
                serie = request.form.get("serie", "").strip()

                if not nome or not turma or not serie:
                    erro = "Nome, turma e série são obrigatórios."
                else:
                    with get_db() as conn:
                        padrao = conn.execute(
                            "SELECT saida_seg, saida_ter, saida_qua, saida_qui, saida_sex FROM horarios_padrao WHERE serie = ?",
                            (serie,)
                        ).fetchone()

                    foto_path = None
                    if 'foto' in request.files:
                        f = request.files['foto']
                        if f and f.filename != '':
                            from app.core.validators import allowed_file
                            from werkzeug.utils import secure_filename
                            if allowed_file(f.filename):
                                fn = secure_filename(f"{datetime.now().strftime('%Y%m%d%H%M%S')}_{f.filename}")
                                f.save(os.path.join(app.config['UPLOAD_FOLDER'], 'photos', fn))
                                foto_path = os.path.join('photos', fn)

                    StudentRepository().create({
                        'nome': nome, 'turma': turma, 'serie': serie,
                        'saida_seg': request.form.get("saida_seg") or (padrao['saida_seg'] if padrao else ''),
                        'saida_ter': request.form.get("saida_ter") or (padrao['saida_ter'] if padrao else ''),
                        'saida_qua': request.form.get("saida_qua") or (padrao['saida_qua'] if padrao else ''),
                        'saida_qui': request.form.get("saida_qui") or (padrao['saida_qui'] if padrao else ''),
                        'saida_sex': request.form.get("saida_sex") or (padrao['saida_sex'] if padrao else ''),
                        'responsaveis': request.form.get("responsaveis", ""),
                        'telefone': '', 'email_responsavel': '',
                        'data_nascimento': '', 'alergias': '', 'observacoes': '',
                        'foto_path': foto_path,
                    })
                    log_operacao(session.get('username'), "CADASTROU ALUNO", f"Nome: {nome}")
                    flash(f"Aluno {nome} cadastrado!", "success")
                    return redirect("/cadastro_aluno")

        return render_template("students/bulk.html", series=Config.SERIES,
                               horarios_padrao=horarios_padrao, erro=erro)

    # ==================== ESQUECI A SENHA ====================
    @app.route("/esqueci_senha", methods=["GET", "POST"])
    def esqueci_senha():
        from app.repositories.users import UserRepository
        from app.core.eamail import enviar_email
        import secrets as _secrets
        from datetime import timedelta

        mensagem = None
        if request.method == "POST":
            email = request.form.get("email", "").strip().lower()
            user = UserRepository().get_by_email(email)
            mensagem = "Se este email estiver cadastrado, você receberá um link em breve."
            if user:
                token = _secrets.token_urlsafe(32)
                expires = (datetime.utcnow() + timedelta(hours=1)).strftime("%Y-%m-%d %H:%M:%S")
                with get_db() as conn:
                    conn.execute("DELETE FROM reset_tokens WHERE user_id = ?", (user['id'],))
                    conn.execute(
                        "INSERT INTO reset_tokens (user_id, token, expires_at) VALUES (?, ?, ?)",
                        (user['id'], token, expires)
                    )
                reset_url = f"{Config.BASE_URL}/redefinir_senha/{token}"
                try:
                    enviar_email(email, "Redefinição de senha — S2E",
                                 f"<p>Clique para redefinir sua senha: <a href='{reset_url}'>{reset_url}</a></p>"
                                 f"<p>Link válido por 1 hora.</p>")
                except Exception:
                    pass

        return render_template("auth/forgot.html", mensagem=mensagem)

    # ==================== REDEFINIR SENHA ====================
    @app.route("/redefinir_senha/<token>", methods=["GET", "POST"])
    def redefinir_senha(token):
        from app.repositories.users import UserRepository

        with get_db() as conn:
            row = conn.execute(
                "SELECT * FROM reset_tokens WHERE token = ? AND expires_at > datetime('now')",
                (token,)
            ).fetchone()

        if not row:
            return render_template("auth/reset.html", token=None, erro="Link inválido ou expirado.")

        erro = None
        if request.method == "POST":
            senha = request.form.get("senha", "")
            confirma = request.form.get("confirma", "")
            if not _senha_forte(senha):
                erro = _MSG_SENHA
            elif senha != confirma:
                erro = "As senhas não coincidem."
            else:
                UserRepository().update_password(dict(row)['user_id'], senha)
                with get_db() as conn:
                    conn.execute("DELETE FROM reset_tokens WHERE token = ?", (token,))
                return redirect("/?resetado=1")

        return render_template("auth/reset.html", token=token, erro=erro)

    return app